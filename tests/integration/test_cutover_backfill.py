"""Phase 2D.1-R5 — identity-aware backfill (docs/ROADMAP.md 2D.1-R5,
docs/PHASE2D1-R0-DECISIONS.md section 4).

Covers the test matrix from the R5 spec section 54: exact-plan rerun,
revised-bytes-same-content rerun, business conflicts, missing/ambiguous
identity, and the no-resurrection bridge invariant. Uses the same
synthetic workbook factories as the rest of the suite
(tests/conftest.py) — no private data.
"""

from __future__ import annotations

import json
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pytest

from bel.adapters.pdf.cmb_bank_statement import ParsedBankTransaction
from bel.application.cutover_backfill import (
    SCOPE_DECISION_CONFIRMATION_DRAFT,
    SCOPE_DECISION_CONFIRMATION_HUMAN,
    SCOPE_DECISION_OUTCOME_OUT_OF_SCOPE,
    SCOPE_DECISION_SOURCE_TYPE,
    _load_payment_scope_decisions,
    backfill_contract_items,
    backfill_contracts,
    backfill_invoices,
    backfill_payment_transactions,
    backfill_payments,
    backfill_procurement_sales_links,
    backfill_sales_contracts,
    backfill_shipments,
)
from bel.application.cutover_plan import CutoverPlanError, run_backfill_plan
from bel.application.sales_contract_facts import create_sales_contract_fact
from bel.domain.contract import Contract
from bel.domain.evidence import EvidenceDocument, EvidenceFragment, FragmentKind
from bel.infrastructure.persistence.database import make_engine, make_session_factory
from bel.infrastructure.persistence.models import Base
from bel.infrastructure.persistence.repositories import (
    ContractRepository,
    EvidenceRepository,
    InvoiceRepository,
    PaymentRepository,
    ProcurementSalesLinkRepository,
    SalesContractRepository,
    ShipmentRepository,
)

NOW = datetime.now(timezone.utc)

CONTRACT_HEADERS = ["序号", "合同编码", "卖方", "买方", "金额", "外销合同编码"]
INVOICE_HEADERS_ROW = [
    "凭证模板", "凭证字号", "发票票种", "开票日期", "发票号码", "数电发票号码", "销方名称",
    "商品名称（明细）", "规格型号（明细）", "单位（明细）", "数量（明细）", "单价（明细）", "金额（明细）",
    "税率（%）（明细）", "税额（明细）", "价税合计（明细）", "发票状态", "发票金额", "发票税额", "发票价税合计",
]


@pytest.fixture
def db_session():
    engine = make_engine("sqlite://")
    Base.metadata.create_all(engine)
    session_factory = make_session_factory(engine)
    with session_factory() as session:
        yield session


def _write_ledger(path: Path, rows: list[list]) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报关出口购销合同"
    ws.append(["Title"])
    ws.append(CONTRACT_HEADERS)
    for row in rows:
        ws.append(row)
    wb.save(path)


def _write_invoices(path: Path, buyer: str, rows: list[list]) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sheet1"
    ws.append([])
    ws.append([buyer])
    ws.append(INVOICE_HEADERS_ROW)
    for row in rows:
        ws.append(row)
    wb.save(path)


def _make_fragment(session):
    doc = EvidenceDocument(
        id=uuid.uuid4(), file_name="x", sha256=uuid.uuid4().hex + uuid.uuid4().hex, source_type="t", imported_at=NOW
    )
    EvidenceRepository(session).add_document(doc)
    frag = EvidenceFragment(
        id=uuid.uuid4(), evidence_document_id=doc.id, fragment_kind=FragmentKind.MANUAL_FACT, sheet_name=None,
        row_number=None, locator_json={}, raw_data={}, created_at=NOW,
    )
    EvidenceRepository(session).add_fragment(frag)
    session.flush()
    return frag


# ---------------------------------------------------------------------------
# A/B/C/D/E — Contract backfill
# ---------------------------------------------------------------------------


def test_a_exact_plan_rerun_no_duplicate(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", "BuyerX", 100]])
    first = backfill_contracts(db_session, path)
    second = backfill_contracts(db_session, path)
    assert first.created == 1
    assert second.created == 0
    assert second.replay_or_corroborating == 0  # is_reimport short-circuit — same bytes
    assert len(ContractRepository(db_session).list_all()) == 1


def test_b_revised_bytes_same_business_content_no_duplicate(db_session, tmp_path):
    path1 = tmp_path / "ledger1.xlsx"
    path2 = tmp_path / "ledger2.xlsx"
    _write_ledger(path1, [[1, "C001", "SellerA", "BuyerX", 100]])
    # Different file bytes (the 序号/sequence column, which is never
    # promoted to a canonical field, differs) but identical business
    # content — this is the R0 bug plain file-sha idempotency cannot
    # solve (spec section 26).
    _write_ledger(path2, [[99, "C001", "SellerA", "BuyerX", 100]])
    first = backfill_contracts(db_session, path1)
    second = backfill_contracts(db_session, path2)
    assert first.created == 1
    assert second.created == 0
    assert second.replay_or_corroborating == 1  # corroborating — different fragment, same content
    assert len(ContractRepository(db_session).list_all()) == 1


def test_c_revised_value_conflict_produces_task_never_overwrites(db_session, tmp_path):
    path1 = tmp_path / "ledger1.xlsx"
    path2 = tmp_path / "ledger2.xlsx"
    _write_ledger(path1, [[1, "C001", "SellerA", "BuyerX", 100]])
    _write_ledger(path2, [[1, "C001", "SellerA", "BuyerX", 999]])  # different amount
    backfill_contracts(db_session, path1)
    second = backfill_contracts(db_session, path2)
    assert len(second.tasks) == 1
    assert second.tasks[0].kind == "CONFLICT"
    current = ContractRepository(db_session).find_by_contract_no("C001")[0]
    assert current.gross_amount == Decimal("100.00")  # untouched — "latest wins" is forbidden


def test_d_missing_identity_produces_task(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", None, "BuyerX", 100]])  # counterparty missing
    result = backfill_contracts(db_session, path)
    assert result.created == 0
    assert result.tasks[0].kind == "IDENTITY_INCOMPLETE"
    assert ContractRepository(db_session).list_all() == []


def test_e_duplicate_identity_ambiguity_never_guesses(db_session, tmp_path):
    frag = _make_fragment(db_session)
    from bel.domain.contract import Contract

    c1 = Contract(
        id=uuid.uuid4(), contract_no="C-DUP", contract_type=None, counterparty="Sup", buyer=None,
        gross_amount=Decimal("1"), currency="CNY", contract_date=None, current_source_fragment_id=frag.id,
        created_at=NOW, updated_at=NOW,
    )
    c2 = Contract(
        id=uuid.uuid4(), contract_no="C-DUP", contract_type=None, counterparty="Sup", buyer=None,
        gross_amount=Decimal("2"), currency="CNY", contract_date=None, current_source_fragment_id=frag.id,
        created_at=NOW, updated_at=NOW,
    )
    ContractRepository(db_session).add(c1)
    ContractRepository(db_session).add(c2)
    db_session.flush()

    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C-DUP", "Sup", "BuyerX", 3]])
    result = backfill_contracts(db_session, path)
    assert result.tasks[0].kind == "IDENTITY_AMBIGUOUS"
    assert len(ContractRepository(db_session).list_all()) == 2  # no merge, no third row


# ---------------------------------------------------------------------------
# F/G — ContractItem backfill
# ---------------------------------------------------------------------------


def test_f_contract_item_exact_replay(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", "BuyerX", 100]])
    backfill_contracts(db_session, path)
    frag = _make_fragment(db_session)
    entries = [{"contract_no": "C001", "counterparty": "SellerA", "source_item_key": "ITEM-1", "fields": {"product_name": "Widget"}}]
    first = backfill_contract_items(db_session, entries, source_fragment_id=frag.id)
    second = backfill_contract_items(db_session, entries, source_fragment_id=frag.id)
    assert first.created == 1
    assert second.replay_or_corroborating == 1


def test_g_contract_item_missing_source_item_key_produces_task(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", "BuyerX", 100]])
    backfill_contracts(db_session, path)
    frag = _make_fragment(db_session)
    entries = [{"contract_no": "C001", "counterparty": "SellerA", "source_item_key": None, "fields": {}}]
    result = backfill_contract_items(db_session, entries, source_fragment_id=frag.id)
    assert result.tasks[0].kind == "IDENTITY_INCOMPLETE"


# ---------------------------------------------------------------------------
# H/I/J — Invoice backfill
# ---------------------------------------------------------------------------


def test_h_invoice_exact_replay_by_external_key(db_session, tmp_path):
    path1 = tmp_path / "inv1.xlsx"
    path2 = tmp_path / "inv2.xlsx"
    row1 = ["T", "V1", "增值税专用发票", date(2026, 1, 5), "INV-1", "DIGITAL-1", "Seller", "Widget", None, "PCS", 10, 10, 100, 0, 0, 100, None, 100, 0, 100]
    # V2 differs (an Evidence-only field, never promoted to a canonical
    # Invoice field) — same real-file-revision scenario as test_b.
    row2 = ["T", "V2", "增值税专用发票", date(2026, 1, 5), "INV-1", "DIGITAL-1", "Seller", "Widget", None, "PCS", 10, 10, 100, 0, 0, 100, None, 100, 0, 100]
    _write_invoices(path1, "BuyerX", [row1])
    _write_invoices(path2, "BuyerX", [row2])
    first = backfill_invoices(db_session, path1, "PURCHASE")
    second = backfill_invoices(db_session, path2, "PURCHASE")
    assert first.created == 1
    assert second.replay_or_corroborating == 1
    assert len(InvoiceRepository(db_session).list_all()) == 1


def test_i_invoice_missing_external_key_produces_task(db_session, tmp_path):
    path = tmp_path / "inv.xlsx"
    row = ["T", "V1", "增值税专用发票", date(2026, 1, 5), "INV-1", None, "Seller", "Widget", None, "PCS", 10, 10, 100, 0, 0, 100, None, 100, 0, 100]
    _write_invoices(path, "BuyerX", [row])
    result = backfill_invoices(db_session, path, "PURCHASE")
    assert result.created == 0
    assert result.tasks[0].kind == "IDENTITY_INCOMPLETE"


def test_j_invoice_same_key_conflicting_content_produces_task(db_session, tmp_path):
    path1 = tmp_path / "inv1.xlsx"
    path2 = tmp_path / "inv2.xlsx"
    row1 = ["T", "V1", "增值税专用发票", date(2026, 1, 5), "INV-1", "DIGITAL-1", "Seller", "Widget", None, "PCS", 10, 10, 100, 0, 0, 100, None, 100, 0, 100]
    row2 = ["T", "V1", "增值税专用发票", date(2026, 1, 5), "INV-1", "DIGITAL-1", "Seller", "Widget", None, "PCS", 10, 10, 200, 0, 0, 200, None, 200, 0, 200]
    _write_invoices(path1, "BuyerX", [row1])
    _write_invoices(path2, "BuyerX", [row2])
    backfill_invoices(db_session, path1, "PURCHASE")
    second = backfill_invoices(db_session, path2, "PURCHASE")
    assert second.tasks[0].kind == "CONFLICT"
    existing = InvoiceRepository(db_session).find_by_external_key("DIGITAL-1")
    assert existing.net_amount == Decimal("100")  # untouched


# ---------------------------------------------------------------------------
# K/L/M/N/O — Payment backfill
# ---------------------------------------------------------------------------


def _txn(index=0, amount=Decimal("100.00"), bank_reference="REF-1", counterparty="Cust"):
    return ParsedBankTransaction(
        page_index=0, transaction_index=index, raw_data={"n": index}, transaction_date=date(2026, 1, 5),
        business_type="转账", bank_reference=bank_reference, description="desc", signed_amount=amount,
        running_balance=Decimal("1000.00"), counterparty=counterparty,
    )


def _seed_scope_contract(session, counterparty="ScopeSupplier", contract_no="SCOPE-1"):
    """A real, Evidence-backed Contract whose counterparty is a party of
    the first-stage procurement scope — the DB state the cutover
    Payment-scope filter resolves against (mirrors the reconciliation
    suite's ``_make_contract``)."""
    doc = EvidenceDocument(
        id=uuid.uuid4(), file_name="scope", sha256=uuid.uuid4().hex + uuid.uuid4().hex, source_type="t",
        imported_at=NOW,
    )
    EvidenceRepository(session).add_document(doc)
    frag = EvidenceFragment(
        id=uuid.uuid4(), evidence_document_id=doc.id, fragment_kind=FragmentKind.MANUAL_FACT, sheet_name=None,
        row_number=None, locator_json={}, raw_data={}, created_at=NOW,
    )
    EvidenceRepository(session).add_fragment(frag)
    session.flush()
    contract = Contract(
        id=uuid.uuid4(), contract_no=contract_no, contract_type="出口报关购销合同", counterparty=counterparty,
        buyer="BuyerX", gross_amount=Decimal("1000.00"), currency="CNY", contract_date=date(2026, 1, 1),
        current_source_fragment_id=frag.id, created_at=NOW, updated_at=NOW,
    )
    ContractRepository(session).add(contract)
    session.flush()
    return contract


def _new_bank_document(session, tag: str):
    doc_id = uuid.uuid4()
    EvidenceRepository(session).add_document(
        EvidenceDocument(
            id=doc_id, file_name=f"bank-{tag}", sha256=uuid.uuid4().hex + uuid.uuid4().hex,
            source_type="cmb_bank_statement_pdf", imported_at=NOW,
        )
    )
    session.flush()
    return doc_id


def _seed_sales_contract(session, *, customer=None, sales_contract_no="SC-1", our_entity="BuyerS"):
    """A genuine, Evidence-backed SalesContract scope (customer optional —
    NULL is a legitimate first-stage state). Mirrors the reconciliation
    suite's sales-contract seeding."""
    frag = _make_fragment(session)
    fields = {"currency": "CNY", "gross_amount": Decimal("500.00"), "contract_date": date(2026, 1, 1)}
    if customer is not None:
        fields["customer"] = customer
    return create_sales_contract_fact(
        session, our_entity=our_entity, sales_contract_no=sales_contract_no, fields=fields,
        source_fragment_id=frag.id, created_at=NOW,
    ).sales_contract


def _confirmed_decision(index, page=0, *, confirmation=SCOPE_DECISION_CONFIRMATION_HUMAN):
    """One synthetic private scope decision anchored to the bank Evidence
    row locator (page, transaction_index)."""
    return {
        (page, index): {
            "decision": SCOPE_DECISION_OUTCOME_OUT_OF_SCOPE,
            "confirmation_type": confirmation,
            "page": page,
            "transaction_index": index,
            "category": "NON_BUSINESS_INBOUND",
            "reason": "synthetic reviewed non-business inbound bank movement",
        }
    }


def test_k_payment_exact_replay(db_session):
    doc_id = uuid.uuid4()
    EvidenceRepository(db_session).add_document(
        EvidenceDocument(id=doc_id, file_name="x", sha256="a" * 64, source_type="cmb_bank_statement_pdf", imported_at=NOW)
    )
    db_session.flush()
    first = backfill_payment_transactions(db_session, [_txn()], source_account_id="ACC-1", document_id=doc_id)
    second = backfill_payment_transactions(db_session, [_txn()], source_account_id="ACC-1", document_id=doc_id)
    assert first.created == 1
    assert second.replay_or_corroborating == 1
    assert len(PaymentRepository(db_session).list_all()) == 1


def test_l_different_source_accounts_are_two_distinct_payments(db_session):
    doc_id = uuid.uuid4()
    EvidenceRepository(db_session).add_document(
        EvidenceDocument(id=doc_id, file_name="x", sha256="b" * 64, source_type="cmb_bank_statement_pdf", imported_at=NOW)
    )
    db_session.flush()
    backfill_payment_transactions(db_session, [_txn()], source_account_id="ACC-1", document_id=doc_id)
    result = backfill_payment_transactions(db_session, [_txn()], source_account_id="ACC-2", document_id=doc_id)
    assert result.created == 1  # same date/direction/amount/bank_reference, DIFFERENT account -> two Payments
    assert len(PaymentRepository(db_session).list_all()) == 2


def test_m_missing_source_account_produces_task_never_silent_dedup(db_session):
    doc_id = uuid.uuid4()
    EvidenceRepository(db_session).add_document(
        EvidenceDocument(id=doc_id, file_name="x", sha256="c" * 64, source_type="cmb_bank_statement_pdf", imported_at=NOW)
    )
    db_session.flush()
    result = backfill_payment_transactions(db_session, [_txn()], source_account_id=None, document_id=doc_id)
    assert result.created == 0
    assert result.tasks[0].kind == "IDENTITY_INCOMPLETE"
    assert result.tasks[0].detail["missing_source_account_id"] is True


def test_n_missing_bank_reference_produces_task(db_session):
    doc_id = uuid.uuid4()
    EvidenceRepository(db_session).add_document(
        EvidenceDocument(id=doc_id, file_name="x", sha256="d" * 64, source_type="cmb_bank_statement_pdf", imported_at=NOW)
    )
    db_session.flush()
    result = backfill_payment_transactions(
        db_session, [_txn(bank_reference=None)], source_account_id="ACC-1", document_id=doc_id
    )
    assert result.tasks[0].detail["missing_bank_reference"] is True


def test_o_same_identity_conflicting_evidence_produces_task(db_session):
    doc_id = uuid.uuid4()
    EvidenceRepository(db_session).add_document(
        EvidenceDocument(id=doc_id, file_name="x", sha256="e" * 64, source_type="cmb_bank_statement_pdf", imported_at=NOW)
    )
    db_session.flush()
    backfill_payment_transactions(db_session, [_txn(counterparty="Cust A")], source_account_id="ACC-1", document_id=doc_id)
    result = backfill_payment_transactions(
        db_session, [_txn(counterparty="Cust B")], source_account_id="ACC-1", document_id=doc_id
    )
    assert result.tasks[0].kind == "CONFLICT"
    assert len(PaymentRepository(db_session).list_all()) == 1  # existing untouched, no second row


# ---------------------------------------------------------------------------
# First-stage Payment-scope filter REPAIR #2 (cutover): the automatic filter
# is procurement-OUT only (M001 membership). Sales-side IN receipts are never
# auto-excluded by negative SalesContract.customer membership; non-business
# IN bank movements are excluded ONLY through an explicit HUMAN_CONFIRMED
# OUT_OF_SCOPE private adjudication anchored to the bank Evidence locator.
# Payment identity is unchanged.
# ---------------------------------------------------------------------------


def test_scope_repair_a_out_unrelated_procurement_counterparty_excluded(db_session):
    _seed_scope_contract(db_session, counterparty="ScopeSupplier")
    doc_id = _new_bank_document(db_session, "A")
    result = backfill_payment_transactions(
        db_session, [_txn(counterparty="UnrelatedOutParty", bank_reference=None, amount=Decimal("-5.00"))],
        source_account_id="ACC-1", document_id=doc_id,
    )
    assert result.created == 0
    assert result.tasks == []
    assert PaymentRepository(db_session).list_all() == []
    assert EvidenceRepository(db_session).find_fragment_by_document(doc_id) is not None


def test_scope_repair_b_out_known_procurement_missing_ref_identity_task(db_session):
    _seed_scope_contract(db_session, counterparty="ScopeSupplier")
    doc_id = _new_bank_document(db_session, "B")
    result = backfill_payment_transactions(
        db_session, [_txn(counterparty="ScopeSupplier", bank_reference=None, amount=Decimal("-100.00"))],
        source_account_id="ACC-1", document_id=doc_id,
    )
    assert result.created == 0
    assert PaymentRepository(db_session).list_all() == []
    assert len(result.tasks) == 1
    assert result.tasks[0].kind == "IDENTITY_INCOMPLETE"
    assert result.tasks[0].detail["missing_bank_reference"] is True


def test_scope_out_known_procurement_complete_identity_payment_created(db_session):
    _seed_scope_contract(db_session, counterparty="ScopeSupplier")
    doc_id = _new_bank_document(db_session, "out-ok")
    result = backfill_payment_transactions(
        db_session, [_txn(counterparty="ScopeSupplier", bank_reference="REF-OUT", amount=Decimal("-100.00"))],
        source_account_id="ACC-1", document_id=doc_id,
    )
    assert result.created == 1
    assert result.tasks == []
    payments = PaymentRepository(db_session).list_all()
    assert len(payments) == 1
    assert payments[0].counterparty == "ScopeSupplier"


def test_scope_repair_c_in_known_sales_customer_not_excluded_no_auto_match(db_session):
    _seed_sales_contract(db_session, customer="KnownCustomer", sales_contract_no="SC-K")
    doc_id = _new_bank_document(db_session, "C")
    result = backfill_payment_transactions(
        db_session, [_txn(counterparty="KnownCustomer", bank_reference="REF-IN-C", amount=Decimal("100.00"))],
        source_account_id="ACC-1", document_id=doc_id,
    )
    assert result.created == 1
    assert result.tasks == []
    # No automatic sales matching of any kind: no MatchCase, no SalesMatchCandidate.
    from bel.infrastructure.persistence.repositories import MatchCaseRepository

    assert MatchCaseRepository(db_session).list_all() == []


def test_scope_repair_d_in_unknown_payer_not_automatically_excluded(db_session):
    """An unknown payer is never automatically OUT_OF_SCOPE, even when a
    populated procurement scope exists (IN is not compared against it)."""
    _seed_scope_contract(db_session, counterparty="ScopeSupplier")
    doc_id = _new_bank_document(db_session, "D")
    result = backfill_payment_transactions(
        db_session, [_txn(counterparty="UnknownPayer", bank_reference="REF-IN-D", amount=Decimal("100.00"))],
        source_account_id="ACC-1", document_id=doc_id,
    )
    assert result.created == 1
    assert result.tasks == []


def test_scope_repair_e_in_sales_customer_null_not_automatically_excluded(db_session):
    _seed_sales_contract(db_session, customer=None, sales_contract_no="SC-NULL")
    doc_id = _new_bank_document(db_session, "E")
    result = backfill_payment_transactions(
        db_session, [_txn(counterparty="SomePayer", bank_reference="REF-IN-E", amount=Decimal("100.00"))],
        source_account_id="ACC-1", document_id=doc_id,
    )
    assert result.created == 1
    assert result.tasks == []


def test_scope_repair_f_in_payer_customer_mismatch_not_automatically_excluded(db_session):
    _seed_sales_contract(db_session, customer="CustomerA", sales_contract_no="SC-A")
    doc_id = _new_bank_document(db_session, "F")
    result = backfill_payment_transactions(
        db_session, [_txn(counterparty="DifferentPayer", bank_reference="REF-IN-F", amount=Decimal("100.00"))],
        source_account_id="ACC-1", document_id=doc_id,
    )
    assert result.created == 1
    assert result.tasks == []


def test_scope_repair_g_in_explicit_human_confirmed_out_of_scope(db_session):
    doc_id = _new_bank_document(db_session, "G")
    result = backfill_payment_transactions(
        db_session, [_txn(counterparty="ReviewedPayer", bank_reference=None, amount=Decimal("100.00"))],
        source_account_id="ACC-1", document_id=doc_id,
        scope_decisions=_confirmed_decision(0),
    )
    assert result.created == 0
    assert result.tasks == []
    assert PaymentRepository(db_session).list_all() == []
    assert EvidenceRepository(db_session).find_fragment_by_document(doc_id) is not None


def test_scope_repair_h_same_in_row_without_adjudication_is_conservative(db_session):
    """Without an explicit adjudication, the SAME IN row is processed
    normally: incomplete identity -> the existing Backfill Task."""
    doc_id = _new_bank_document(db_session, "H")
    result = backfill_payment_transactions(
        db_session, [_txn(counterparty="ReviewedPayer", bank_reference=None, amount=Decimal("100.00"))],
        source_account_id="ACC-1", document_id=doc_id,
    )
    assert result.created == 0
    assert len(result.tasks) == 1
    assert result.tasks[0].kind == "IDENTITY_INCOMPLETE"


def test_scope_repair_i_missing_bank_reference_alone_never_excludes(db_session):
    """Missing bank_reference alone never triggers OUT_OF_SCOPE: with a
    populated scope present, an IN row missing its reference still flows
    into frozen identity validation (Task), never into an exclusion."""
    _seed_scope_contract(db_session, counterparty="ScopeSupplier")
    doc_id = _new_bank_document(db_session, "I")
    result = backfill_payment_transactions(
        db_session, [_txn(counterparty="RandomPayer", bank_reference=None, amount=Decimal("100.00"))],
        source_account_id="ACC-1", document_id=doc_id,
    )
    assert result.created == 0
    assert len(result.tasks) == 1
    assert result.tasks[0].kind == "IDENTITY_INCOMPLETE"
    assert result.tasks[0].detail["missing_bank_reference"] is True


def test_scope_repair_j_scope_locator_not_usable_as_payment_identity(db_session):
    """The adjudication locator pinpoints a SOURCE Evidence row; it is not
    Payment business identity. Two rows with identical business content at
    DIFFERENT locators: adjudicating locator (0,0) excludes only that row —
    the identical-content row at (0,1) still creates a normal Payment and
    neither is affected by a dedup/conflict via the locator."""
    doc_id = _new_bank_document(db_session, "J")
    t0 = _txn(index=0, counterparty="SameInPayer", bank_reference="REF-J", amount=Decimal("50.00"))
    t1 = _txn(index=1, counterparty="SameInPayer", bank_reference="REF-J", amount=Decimal("50.00"))
    result = backfill_payment_transactions(
        db_session, [t0, t1], source_account_id="ACC-1", document_id=doc_id,
        scope_decisions=_confirmed_decision(0),
    )
    assert result.created == 1
    assert result.tasks == []
    payments = PaymentRepository(db_session).list_all()
    assert len(payments) == 1
    assert payments[0].bank_reference == "REF-J"


def test_scope_decision_loader_applies_confirmed_and_evidences_but_not_draft(tmp_path, db_session):
    """File adjudication: HUMAN_CONFIRMED OUT_OF_SCOPE is applied AND
    recorded as Evidence (source_type cutover_payment_scope_decision); a
    DRAFT_NEEDS_BUSINESS_CONFIRMATION entry is never applied."""
    from bel.adapters.common import compute_sha256

    txns = [
        _txn(index=0, counterparty="Any", bank_reference="R0", amount=Decimal("10.00")),
        _txn(index=1, counterparty="Any", bank_reference="R1", amount=Decimal("10.00")),
    ]
    source_sha = "a" * 64
    artifact = {
        "version": 1,
        "source_sha256": source_sha,
        "entries": [
            {
                "decision": SCOPE_DECISION_OUTCOME_OUT_OF_SCOPE,
                "confirmation_type": SCOPE_DECISION_CONFIRMATION_HUMAN,
                "page": 0, "transaction_index": 0,
                "category": "NON_BUSINESS_INBOUND", "reason": "synthetic confirmed",
            },
            {
                "decision": SCOPE_DECISION_OUTCOME_OUT_OF_SCOPE,
                "confirmation_type": SCOPE_DECISION_CONFIRMATION_DRAFT,
                "page": 0, "transaction_index": 1,
                "category": "NON_BUSINESS_INBOUND", "reason": "synthetic draft",
            },
        ]
    }
    path = tmp_path / "payment-scope-decisions.json"
    path.write_text(json.dumps(artifact), encoding="utf-8")
    applied = _load_payment_scope_decisions(
        db_session, path, txns, expected_source_sha256=source_sha, now=NOW
    )
    assert (0, 0) in applied
    assert (0, 1) not in applied  # DRAFT is never applied

    doc = EvidenceRepository(db_session).find_document_by_sha256(compute_sha256(path))
    assert doc is not None
    assert doc.source_type == SCOPE_DECISION_SOURCE_TYPE
    assert EvidenceRepository(db_session).find_fragment_by_document(doc.id) is not None


def test_plan_rejects_scope_decisions_under_expected(tmp_path, db_session):
    period_dir = tmp_path / "period"
    period_dir.mkdir()
    plan = {
        "version": 1,
        "payments": [
            {"path": "bank/stmt.pdf", "profile": "cmb", "source_account_id": "ACC",
             "scope_decisions": "expected/payment-scope-decisions.json"},
        ],
    }
    with pytest.raises(CutoverPlanError):
        run_backfill_plan(db_session, plan, period_dir=period_dir, created_at=NOW)


def test_plan_rejects_unknown_payments_entry_key(tmp_path, db_session):
    period_dir = tmp_path / "period"
    period_dir.mkdir()
    plan = {
        "version": 1,
        "payments": [{"path": "bank/stmt.pdf", "profile": "cmb", "source_account_id": "ACC", "surprise": True}],
    }
    with pytest.raises(CutoverPlanError):
        run_backfill_plan(db_session, plan, period_dir=period_dir, created_at=NOW)


def test_plan_rejects_scope_decisions_path_escape(tmp_path, db_session):
    period_dir = tmp_path / "period"
    period_dir.mkdir()
    plan = {
        "version": 1,
        "payments": [
            {"path": "bank/stmt.pdf", "profile": "cmb", "source_account_id": "ACC",
             "scope_decisions": "../payment-scope-decisions.json"},
        ],
    }
    with pytest.raises(CutoverPlanError):
        run_backfill_plan(db_session, plan, period_dir=period_dir, created_at=NOW)


def test_plan_wires_scope_decisions_end_to_end(tmp_path, db_session):
    """The plan references a private adjudication artifact under facts/,
    the wrapper applies a HUMAN_CONFIRMED OUT_OF_SCOPE decision (an
    in-scope OUT row excluded by adjudication), and the adjudication is
    recorded as Evidence — plan stays orchestration."""
    from bel.adapters.common import compute_sha256
    from fixtures.synthetic.bank_pdf import build_cmb_bank_statement_pdf

    period_dir = tmp_path / "period"
    (period_dir / "contracts").mkdir(parents=True)
    (period_dir / "bank").mkdir()
    (period_dir / "facts").mkdir()

    _write_ledger(period_dir / "contracts" / "ledger.xlsx", [[1, "SCOPE-C", "ScopeSupplier", "BuyerX", 100]])
    pdf = period_dir / "bank" / "stmt.pdf"
    build_cmb_bank_statement_pdf(
        pdf, "1000.00", [("20260105", "对公转账出", "REF-A", "desc", "ScopeSupplier", "100.00")]
    )
    decisions_file = period_dir / "facts" / "payment-scope-decisions.json"
    decisions_file.write_text(
        json.dumps(
            {
                "version": 1,
                "source_sha256": compute_sha256(pdf),
                "entries": [
                    {
                        "decision": SCOPE_DECISION_OUTCOME_OUT_OF_SCOPE,
                        "confirmation_type": SCOPE_DECISION_CONFIRMATION_HUMAN,
                        "page": 0, "transaction_index": 0,
                        "category": "NON_BUSINESS_INBOUND", "reason": "plan plumbing synthetic",
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    plan = {
        "version": 1,
        "contracts": {"path": "contracts/ledger.xlsx"},
        "payments": [
            {"path": "bank/stmt.pdf", "profile": "cmb", "source_account_id": "ACC-P",
             "scope_decisions": "facts/payment-scope-decisions.json"},
        ],
    }
    result = run_backfill_plan(db_session, plan, period_dir=period_dir, created_at=NOW)
    assert result.sections["contracts"]["created"] == 1
    payments = result.sections["payments"][0]
    assert payments["created"] == 0
    assert payments["tasks"] == []
    assert PaymentRepository(db_session).list_all() == []
    doc = EvidenceRepository(db_session).find_document_by_sha256(compute_sha256(decisions_file))
    assert doc is not None
    assert doc.source_type == SCOPE_DECISION_SOURCE_TYPE


# ---------------------------------------------------------------------------
# Repair #3 — scope adjudication source-SHA binding, strict/unique locator
# schema, exact one-to-one resolution. No private values anywhere.
# ---------------------------------------------------------------------------


def _dec_entry(page=0, index=0, confirmation=SCOPE_DECISION_CONFIRMATION_HUMAN):
    return {
        "decision": SCOPE_DECISION_OUTCOME_OUT_OF_SCOPE,
        "confirmation_type": confirmation,
        "page": page, "transaction_index": index,
        "category": "NON_BUSINESS_INBOUND", "reason": "synthetic reviewed decision",
    }


def _write_scope_artifact(tmp_path, entries, *, source_sha="a" * 64, name="scope.json", include_sha=True):
    payload = {"version": 1}
    if include_sha:
        payload["source_sha256"] = source_sha
    payload["entries"] = entries
    path = tmp_path / name
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def _load_scope(path, db_session, txns, expected="a" * 64):
    return _load_payment_scope_decisions(
        db_session, path, txns, expected_source_sha256=expected, now=NOW
    )


def _two_txn_rows():
    return [_txn(index=0), _txn(index=1)]


def test_repair3_wrong_source_sha_hard_fail(tmp_path, db_session):
    art = _write_scope_artifact(tmp_path, [_dec_entry(0, 0)], source_sha="b" * 64)
    with pytest.raises(ValueError):
        _load_scope(art, db_session, _two_txn_rows(), expected="a" * 64)


def test_repair3_missing_source_sha256_hard_fail(tmp_path, db_session):
    art = _write_scope_artifact(tmp_path, [_dec_entry(0, 0)], include_sha=False)
    with pytest.raises(ValueError):
        _load_scope(art, db_session, _two_txn_rows())


def test_repair3_unsupported_artifact_version_hard_fail(tmp_path, db_session):
    art = _write_scope_artifact(tmp_path, [_dec_entry(0, 0)])
    payload = json.loads(art.read_text(encoding="utf-8"))
    payload["version"] = 2
    art.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(ValueError):
        _load_scope(art, db_session, _two_txn_rows())


def test_repair3_malformed_source_sha256_hard_fail(tmp_path, db_session):
    art = _write_scope_artifact(tmp_path, [_dec_entry(0, 0)], source_sha="not-a-sha256")
    with pytest.raises(ValueError):
        _load_scope(art, db_session, _two_txn_rows())


def test_repair3_page_numeric_string_rejected(tmp_path, db_session):
    art = _write_scope_artifact(tmp_path, [_dec_entry(page="0", index=0)])
    with pytest.raises(ValueError):
        _load_scope(art, db_session, _two_txn_rows())


def test_repair3_transaction_index_numeric_string_rejected(tmp_path, db_session):
    art = _write_scope_artifact(tmp_path, [_dec_entry(page=0, index="4")])
    with pytest.raises(ValueError):
        _load_scope(art, db_session, _two_txn_rows())


def test_repair3_page_boolean_rejected(tmp_path, db_session):
    art = _write_scope_artifact(tmp_path, [dict(_dec_entry(0, 0), page=True)])
    with pytest.raises(ValueError):
        _load_scope(art, db_session, _two_txn_rows())


def test_repair3_float_locator_rejected(tmp_path, db_session):
    art = _write_scope_artifact(tmp_path, [dict(_dec_entry(0, 0), page=0.0)])
    with pytest.raises(ValueError):
        _load_scope(art, db_session, _two_txn_rows())


def test_repair3_negative_locator_rejected(tmp_path, db_session):
    art = _write_scope_artifact(tmp_path, [_dec_entry(0, -1)])
    with pytest.raises(ValueError):
        _load_scope(art, db_session, _two_txn_rows())


def test_repair3_duplicate_locator_rejected(tmp_path, db_session):
    # Duplicates are rejected even when the payload is byte-identical.
    art = _write_scope_artifact(tmp_path, [_dec_entry(0, 0), _dec_entry(0, 0)])
    with pytest.raises(ValueError):
        _load_scope(art, db_session, _two_txn_rows())


def test_repair3_zero_match_locator_rejected(tmp_path, db_session):
    art = _write_scope_artifact(tmp_path, [_dec_entry(0, 9)])  # only (0,0),(0,1) exist
    with pytest.raises(ValueError):
        _load_scope(art, db_session, _two_txn_rows())


def test_repair3_draft_only_not_applied_and_not_evidenced(tmp_path, db_session):
    from bel.adapters.common import compute_sha256

    art = _write_scope_artifact(tmp_path, [_dec_entry(0, 0, confirmation=SCOPE_DECISION_CONFIRMATION_DRAFT)])
    applied = _load_scope(art, db_session, _two_txn_rows())
    assert applied == {}
    assert EvidenceRepository(db_session).find_document_by_sha256(compute_sha256(art)) is None


def test_repair3_unique_valid_locator_applies(tmp_path, db_session):
    from bel.adapters.common import compute_sha256

    art = _write_scope_artifact(tmp_path, [_dec_entry(0, 0), _dec_entry(0, 1, confirmation=SCOPE_DECISION_CONFIRMATION_DRAFT)])
    applied = _load_scope(art, db_session, _two_txn_rows())
    assert set(applied) == {(0, 0)}
    assert EvidenceRepository(db_session).find_document_by_sha256(compute_sha256(art)) is not None


def _wrapper_case(tmp_path, db_session, rows, *, scope_entries, seed_contract=True):
    """Build a synthetic CMB PDF + a source-bound adjudication artifact and
    run the cutover wrapper against them (single, fresh bank source)."""
    from bel.adapters.common import compute_sha256
    from fixtures.synthetic.bank_pdf import build_cmb_bank_statement_pdf

    period = tmp_path / "period"
    (period / "bank").mkdir(parents=True)
    (period / "facts").mkdir()
    if seed_contract:
        _seed_scope_contract(db_session, counterparty="ScopeSupplier")
    pdf = period / "bank" / "stmt.pdf"
    build_cmb_bank_statement_pdf(pdf, "1000.00", rows)
    art = period / "facts" / "payment-scope-decisions.json"
    art.write_text(
        json.dumps({"version": 1, "source_sha256": compute_sha256(pdf), "entries": scope_entries}),
        encoding="utf-8",
    )
    return pdf, art


def test_repair3_valid_human_confirmed_wrapper_no_payment_no_task_evidence_retained(tmp_path, db_session):
    pdf, art = _wrapper_case(
        tmp_path, db_session,
        [("20260105", "对公转账出", "REF-A", "desc", "ScopeSupplier", "10.00")],
        scope_entries=[_dec_entry(0, 0)],
    )
    result = backfill_payments(db_session, pdf, "cmb", source_account_id="ACC-1", scope_decisions_path=art)
    assert result.created == 0
    assert result.tasks == []
    assert PaymentRepository(db_session).list_all() == []
    # Bank EvidenceFragment for the excluded row is retained.
    from bel.adapters.common import compute_sha256

    bank_doc = EvidenceRepository(db_session).find_document_by_sha256(compute_sha256(pdf))
    assert bank_doc is not None
    assert EvidenceRepository(db_session).find_fragment_by_document(bank_doc.id) is not None
    # Adjudication recorded as Evidence.
    decision_doc = EvidenceRepository(db_session).find_document_by_sha256(compute_sha256(art))
    assert decision_doc is not None
    assert decision_doc.source_type == SCOPE_DECISION_SOURCE_TYPE


def test_repair3_cross_source_artifact_reuse_rejected(tmp_path, db_session):
    """One scope artifact cannot bind to two different payment sources: a
    mismatch hard-fails BEFORE any bank EvidenceDocument is created."""
    from bel.adapters.common import compute_sha256
    from fixtures.synthetic.bank_pdf import build_cmb_bank_statement_pdf

    _seed_scope_contract(db_session, counterparty="ScopeSupplier")
    pdf_a = tmp_path / "a.pdf"
    build_cmb_bank_statement_pdf(pdf_a, "1000.00", [("20260105", "对公转账出", "REF-A", "d", "ScopeSupplier", "10.00")])
    sha_a = compute_sha256(pdf_a)
    pdf_b = tmp_path / "b.pdf"
    build_cmb_bank_statement_pdf(pdf_b, "2000.00", [("20260106", "对公转账出", "REF-B", "d", "ScopeSupplier", "30.00")])
    sha_b = compute_sha256(pdf_b)
    art = tmp_path / "scope.json"
    art.write_text(json.dumps({"version": 1, "source_sha256": sha_a, "entries": [_dec_entry(0, 0)]}), encoding="utf-8")

    # Reusing the artifact confirmed for source A against source B -> hard fail.
    with pytest.raises(ValueError):
        backfill_payments(db_session, pdf_b, "cmb", source_account_id="ACC-1", scope_decisions_path=art)
    # Nothing was partially persisted for B, so a corrected retry is never masked.
    assert EvidenceRepository(db_session).find_document_by_sha256(sha_b) is None
    assert PaymentRepository(db_session).list_all() == []

    # The same artifact correctly applies to its own source A.
    result_a = backfill_payments(db_session, pdf_a, "cmb", source_account_id="ACC-1", scope_decisions_path=art)
    assert result_a.created == 0
    assert result_a.tasks == []


def test_repair3_source_sha256_never_becomes_payment_identity(tmp_path, db_session):
    """source_sha256 is provenance binding only. It never appears on a
    Payment, never in an identity-incomplete Task key, and never in the
    per-entry adjudication Evidence."""
    from bel.adapters.common import compute_sha256

    pdf, art = _wrapper_case(
        tmp_path, db_session,
        [("20260105", "对公转账出", "REF-A", "desc", "ScopeSupplier", "10.00"),
         ("20260106", "对公转账出", "REF-B", "desc", "ScopeSupplier", "20.00")],
        scope_entries=[_dec_entry(0, 0)],
    )
    sha = compute_sha256(pdf)
    result = backfill_payments(db_session, pdf, "cmb", source_account_id="ACC-1", scope_decisions_path=art)
    # Row (0,0) excluded; in-scope row (0,1) still flows to frozen identity
    # validation (the synthetic PDF carries no parsed reference -> Task).
    assert result.created == 0
    assert len(result.tasks) == 1
    assert result.tasks[0].kind == "IDENTITY_INCOMPLETE"
    for task in result.tasks:
        assert sha not in task.detail["identity_key"]
    assert all(not hasattr(p, "source_sha256") for p in PaymentRepository(db_session).list_all())
    # The adjudication Evidence fragment carries only the decision fields.
    decision_doc = EvidenceRepository(db_session).find_document_by_sha256(compute_sha256(art))
    frag = EvidenceRepository(db_session).find_fragment_by_document(decision_doc.id)
    assert "source_sha256" not in frag.raw_data


# ---------------------------------------------------------------------------
# P/Q — Shipment backfill (delegates to R2's create_shipment_fact)
# ---------------------------------------------------------------------------


def test_p_shipment_complete_identity_replay(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", "BuyerX", 100]])
    backfill_contracts(db_session, path)
    frag = _make_fragment(db_session)
    entries = [{"contract_no": "C001", "counterparty": "SellerA", "external_reference": "EXP-1", "execution_date": date(2026, 1, 10), "quantity": Decimal("5")}]
    first = backfill_shipments(db_session, entries, source_fragment_id=frag.id)
    second = backfill_shipments(db_session, entries, source_fragment_id=frag.id)
    assert first.created == 1
    assert second.replay_or_corroborating == 1


def test_q_shipment_incomplete_identity_never_silently_deduped(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", "BuyerX", 100]])
    backfill_contracts(db_session, path)
    frag = _make_fragment(db_session)
    entries = [{"contract_no": "C001", "counterparty": "SellerA", "external_reference": None, "execution_date": date(2026, 1, 10), "quantity": Decimal("5")}]
    result = backfill_shipments(db_session, entries, source_fragment_id=frag.id)
    assert result.created == 0
    assert result.tasks[0].kind == "CONFLICT"  # ShipmentIdentityIncomplete surfaces as a Task, no Shipment created
    assert ShipmentRepository(db_session).list_all() == []


# ---------------------------------------------------------------------------
# R/S/T — SalesContract backfill
# ---------------------------------------------------------------------------


def test_r_sales_contract_missing_our_entity_rejected(db_session):
    frag = _make_fragment(db_session)
    entries = [{"our_entity": None, "sales_contract_no": "SC-1", "fields": {}}]
    result = backfill_sales_contracts(db_session, entries, source_fragment_id=frag.id)
    assert result.tasks[0].kind == "CONFLICT"
    assert SalesContractRepository(db_session).list_all() == []


def test_s_sales_contract_missing_sales_contract_no_rejected(db_session):
    frag = _make_fragment(db_session)
    entries = [{"our_entity": "Entity A", "sales_contract_no": None, "fields": {}}]
    result = backfill_sales_contracts(db_session, entries, source_fragment_id=frag.id)
    assert result.tasks[0].kind == "CONFLICT"


def test_t_sales_contract_customer_null_still_creates_anchor(db_session):
    frag = _make_fragment(db_session)
    entries = [{"our_entity": "Entity A", "sales_contract_no": "SC-1", "fields": {}}]
    result = backfill_sales_contracts(db_session, entries, source_fragment_id=frag.id)
    assert result.created == 1
    sc = SalesContractRepository(db_session).list_all()[0]
    assert sc.customer is None


# ---------------------------------------------------------------------------
# V/W/X — ProcurementSalesLink backfill (replay, no-resurrection, reestablish)
# ---------------------------------------------------------------------------


def test_v_link_current_replay(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", "BuyerX", 100]])
    backfill_contracts(db_session, path)
    frag = _make_fragment(db_session)
    backfill_sales_contracts(db_session, [{"our_entity": "Entity A", "sales_contract_no": "SC-1", "fields": {}}], source_fragment_id=frag.id)

    entries = [{"contract_no": "C001", "counterparty": "SellerA", "sales_our_entity": "Entity A", "sales_contract_no": "SC-1"}]
    first = backfill_procurement_sales_links(db_session, entries, source_fragment_id=frag.id)
    second = backfill_procurement_sales_links(db_session, entries, source_fragment_id=frag.id)
    assert first.created == 1
    assert second.replay_or_corroborating == 1


def test_w_historical_link_replay_never_resurrects(db_session, tmp_path):
    """Create P->S episode E1, invalidate it, then rerun backfill with
    the SAME historical Evidence — the retired episode must NOT become
    current again (section 16/46, HARD)."""
    from bel.application.procurement_sales_link import correct_procurement_sales_link
    from bel.domain.procurement_sales_link import ConfirmationType as LinkConfirmationType

    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", "BuyerX", 100]])
    backfill_contracts(db_session, path)
    frag = _make_fragment(db_session)
    backfill_sales_contracts(db_session, [{"our_entity": "Entity A", "sales_contract_no": "SC-1", "fields": {}}], source_fragment_id=frag.id)

    entries = [{"contract_no": "C001", "counterparty": "SellerA", "sales_our_entity": "Entity A", "sales_contract_no": "SC-1"}]
    first = backfill_procurement_sales_links(db_session, entries, source_fragment_id=frag.id)
    assert first.created == 1

    contract = ContractRepository(db_session).find_by_contract_no("C001")[0]
    sales_contract = SalesContractRepository(db_session).list_all()[0]
    link = ProcurementSalesLinkRepository(db_session).get_current_link(contract.id, sales_contract.id)

    invalidate_frag = _make_fragment(db_session)
    correct_procurement_sales_link(
        db_session, superseded_link_id=link.id, source_fragment_id=invalidate_frag.id,
        confirmation_type=LinkConfirmationType.HUMAN_CONFIRMED, created_at=NOW,
    )

    # Rerun with the SAME (frag.id) historical Evidence.
    rerun = backfill_procurement_sales_links(db_session, entries, source_fragment_id=frag.id)
    assert ProcurementSalesLinkRepository(db_session).get_current_link(contract.id, sales_contract.id) is None
    # add_procurement_sales_link never resurrects — a replay against a
    # retired-only business key is either a no-op or an explicit new
    # ADD attempt rejected by the link module's own semantics; either
    # way, no current episode may result from this rerun.


def test_x_explicit_reestablish_still_works_outside_backfill(db_session, tmp_path):
    from bel.application.procurement_sales_link import (
        correct_procurement_sales_link,
        reestablish_procurement_sales_link,
    )
    from bel.domain.procurement_sales_link import ConfirmationType as LinkConfirmationType

    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", "BuyerX", 100]])
    backfill_contracts(db_session, path)
    frag = _make_fragment(db_session)
    backfill_sales_contracts(db_session, [{"our_entity": "Entity A", "sales_contract_no": "SC-1", "fields": {}}], source_fragment_id=frag.id)
    entries = [{"contract_no": "C001", "counterparty": "SellerA", "sales_our_entity": "Entity A", "sales_contract_no": "SC-1"}]
    backfill_procurement_sales_links(db_session, entries, source_fragment_id=frag.id)

    contract = ContractRepository(db_session).find_by_contract_no("C001")[0]
    sales_contract = SalesContractRepository(db_session).list_all()[0]
    link = ProcurementSalesLinkRepository(db_session).get_current_link(contract.id, sales_contract.id)
    invalidate_frag = _make_fragment(db_session)
    correct_procurement_sales_link(
        db_session, superseded_link_id=link.id, source_fragment_id=invalidate_frag.id,
        confirmation_type=LinkConfirmationType.HUMAN_CONFIRMED, created_at=NOW,
    )

    new_frag = _make_fragment(db_session)
    reestablish_procurement_sales_link(
        db_session, procurement_contract_id=contract.id, sales_contract_id=sales_contract.id,
        source_fragment_id=new_frag.id, confirmation_type=LinkConfirmationType.HUMAN_CONFIRMED, created_at=NOW,
    )
    assert ProcurementSalesLinkRepository(db_session).get_current_link(contract.id, sales_contract.id) is not None


# ---------------------------------------------------------------------------
# 47 — backfill never reads expected/ as a Fact source (structural check)
# ---------------------------------------------------------------------------


def test_backfill_module_never_touches_private_root_resolution():
    import inspect

    import bel.application.cutover_backfill as module

    source = inspect.getsource(module)
    assert "os.environ" not in source
    assert "resolve_private_root" not in source
    assert "import os" not in source


# ---------------------------------------------------------------------------
# Gate-fix — frozen legacy-ledger sales-scope basis (section 2)
# ---------------------------------------------------------------------------


def test_sales_scope_basis_established_from_same_row(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", "BuyerX", 100, "SC-BASIS-1"]])
    backfill_contracts(db_session, path)

    contract = ContractRepository(db_session).find_by_contract_no("C001")[0]
    sales_contracts = SalesContractRepository(db_session).list_all()
    assert len(sales_contracts) == 1
    sc = sales_contracts[0]
    assert sc.our_entity == "BuyerX"  # 买方 on the SAME row
    assert sc.sales_contract_no == "SC-BASIS-1"  # 外销合同编码 on the SAME row
    assert sc.customer is None  # never inferred
    assert ProcurementSalesLinkRepository(db_session).get_current_link(contract.id, sc.id) is not None


def test_sales_scope_basis_absent_when_export_code_missing(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", "BuyerX", 100, None]])
    backfill_contracts(db_session, path)
    assert SalesContractRepository(db_session).list_all() == []


def test_sales_scope_basis_absent_when_buyer_missing(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", None, 100, "SC-BASIS-2"]])
    backfill_contracts(db_session, path)
    assert SalesContractRepository(db_session).list_all() == []


def test_sales_scope_basis_never_stitches_across_rows(db_session, tmp_path):
    """Row 1 supplies our_entity (买方) but no export code; row 2 supplies
    an export code but has a DIFFERENT buyer — neither row alone
    satisfies both halves, and the two must never be combined."""
    path = tmp_path / "ledger.xlsx"
    _write_ledger(
        path,
        [
            [1, "C001", "SellerA", "BuyerX", 100, None],
            [2, "C002", "SellerB", "BuyerY", 200, "SC-BASIS-3"],
        ],
    )
    backfill_contracts(db_session, path)
    sales_contracts = SalesContractRepository(db_session).list_all()
    assert len(sales_contracts) == 1
    assert sales_contracts[0].our_entity == "BuyerY"  # only row 2's OWN pair, never row 1's buyer


def test_sales_scope_basis_reruns_idempotently(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", "SellerA", "BuyerX", 100, "SC-BASIS-4"]])
    backfill_contracts(db_session, path)
    path2 = tmp_path / "ledger2.xlsx"
    _write_ledger(path2, [[99, "C001", "SellerA", "BuyerX", 100, "SC-BASIS-4"]])
    backfill_contracts(db_session, path2)  # revised bytes, same business content
    assert len(SalesContractRepository(db_session).list_all()) == 1


# ---------------------------------------------------------------------------
# Gate-fix — persistent, idempotent OPEN TaskException (section 1)
# ---------------------------------------------------------------------------


def test_identity_incomplete_task_is_persisted_open(db_session, tmp_path):
    path = tmp_path / "ledger.xlsx"
    _write_ledger(path, [[1, "C001", None, "BuyerX", 100, None]])
    result = backfill_contracts(db_session, path)
    assert len(result.tasks) == 1
    task_id = result.tasks[0].task_exception_id

    from bel.domain.exception import ExceptionStatus, ExceptionType
    from bel.infrastructure.persistence.repositories import ExceptionRepository

    persisted = [e for e in ExceptionRepository(db_session).list_all() if e.id == task_id]
    assert len(persisted) == 1
    assert persisted[0].status == ExceptionStatus.OPEN
    assert persisted[0].exception_type == ExceptionType.BACKFILL_IDENTITY_INCOMPLETE
    assert persisted[0].detail["identity_key"]  # structured, not a parsed summary


def test_identity_incomplete_task_rerun_reuses_same_open_task(db_session, tmp_path):
    """A rerun hitting the SAME underlying identity problem must reuse
    the existing OPEN TaskException, never create a duplicate."""
    path1 = tmp_path / "ledger1.xlsx"
    path2 = tmp_path / "ledger2.xlsx"
    _write_ledger(path1, [[1, "C001", None, "BuyerX", 100, None]])
    _write_ledger(path2, [[99, "C001", None, "BuyerX", 100, None]])  # revised bytes, same problem

    from bel.domain.exception import ExceptionType
    from bel.infrastructure.persistence.repositories import ExceptionRepository

    result1 = backfill_contracts(db_session, path1)
    result2 = backfill_contracts(db_session, path2)
    assert result1.tasks[0].task_exception_id == result2.tasks[0].task_exception_id
    open_tasks = [
        e for e in ExceptionRepository(db_session).list_open() if e.exception_type == ExceptionType.BACKFILL_IDENTITY_INCOMPLETE
    ]
    assert len(open_tasks) == 1


def test_conflict_task_is_persisted_open(db_session, tmp_path):
    path1 = tmp_path / "ledger1.xlsx"
    path2 = tmp_path / "ledger2.xlsx"
    _write_ledger(path1, [[1, "C001", "SellerA", "BuyerX", 100, None]])
    _write_ledger(path2, [[1, "C001", "SellerA", "BuyerX", 999, None]])
    backfill_contracts(db_session, path1)
    result = backfill_contracts(db_session, path2)

    from bel.domain.exception import ExceptionStatus, ExceptionType
    from bel.infrastructure.persistence.repositories import ExceptionRepository

    task_id = result.tasks[0].task_exception_id
    persisted = ExceptionRepository(db_session).list_all()
    match = next(e for e in persisted if e.id == task_id)
    assert match.status == ExceptionStatus.OPEN
    assert match.exception_type == ExceptionType.BACKFILL_CONFLICT
